"""Tests for building a deck from a notebook.

Most of these are regressions. The markdown rules and the layout guards all exist because a
generated deck fails *silently* - a swallowed bullet or a table running off the slide looks fine in
the XML and is only visible once the deck is rendered - so each rule gets a test rather than a
comment.
"""

import base64
import io

import pytest

nbformat = pytest.importorskip("nbformat")
pytest.importorskip("pptx")
pytest.importorskip("PIL")
pytest.importorskip("lxml")
pd = pytest.importorskip("pandas")

from portmanteau.slides.notebook_deck import (  # noqa: E402
    DeckError, DeckStyle, build_deck, estimate_height, inline_runs, parse_bullets, read_notebook)


# --------------------------------------------------------------------------- helpers

def png_bytes(width=40, height=30):
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (width, height), "white").save(buffer, format="PNG")
    return base64.b64encode(buffer.getvalue()).decode()


def code(source, outputs=(), execution_count=1):
    cell = nbformat.v4.new_code_cell(source)
    cell.outputs = list(outputs)
    cell.execution_count = execution_count
    return cell


def display(data):
    return nbformat.v4.new_output("display_data", data=data)


def frame_html(frame):
    return {"text/html": frame.to_html(), "text/plain": repr(frame)}


def notebook(cells):
    nb = nbformat.v4.new_notebook()
    nb.cells = cells
    return nb


def write(tmp_path, nb, name="deck.ipynb"):
    path = tmp_path / name
    nbformat.write(nb, str(path))
    return path


def simple_notebook(**kwargs):
    """A one-slide notebook: title, text, and whatever images are passed."""
    images = kwargs.get("images")
    cells = [nbformat.v4.new_markdown_cell("# Deck title\n\nExplanatory preamble."),
             code("import pandas", execution_count=1),
             nbformat.v4.new_markdown_cell("## First slide"),
             nbformat.v4.new_markdown_cell("### text\n\n- one bullet\n- another")]
    if images is not None:
        cells.append(nbformat.v4.new_markdown_cell("### images"))
        cells.append(images)
    cells.append(nbformat.v4.new_markdown_cell("### notes\n\nSpeaker notes."))
    return notebook(cells)


# --------------------------------------------------------------------------- parsing

def test_reads_title_sections_and_ignores_preamble(tmp_path):
    path = write(tmp_path, simple_notebook(images=code("fig", [display({"image/png": png_bytes()})])))
    title, slides = read_notebook(path)
    assert title == "Deck title"                      # the `# ` line only, not the whole cell
    assert [s.title for s in slides] == ["First slide"]
    assert len(slides[0].cells["text"]) == 1
    assert len(slides[0].cells["images"]) == 1
    assert slides[0].cells["prep"] == []               # the preamble code cell is not a section


def test_unknown_section_is_refused(tmp_path):
    nb = simple_notebook(images=None)
    nb.cells[3].source = "### bullets\n\n- one"
    with pytest.raises(DeckError, match="unknown section"):
        read_notebook(write(tmp_path, nb))


def test_missing_deck_title_is_refused(tmp_path):
    nb = simple_notebook(images=None)
    nb.cells[0].source = "Not a heading"
    with pytest.raises(DeckError, match="no '# ' heading"):
        read_notebook(write(tmp_path, nb))


def test_unexecuted_and_errored_cells_are_refused(tmp_path):
    nb = simple_notebook(images=code("fig", [display({"image/png": png_bytes()})],
                                     execution_count=None))
    with pytest.raises(DeckError, match="never executed"):
        read_notebook(write(tmp_path, nb))

    error = nbformat.v4.new_output("error", ename="ValueError", evalue="boom", traceback=[])
    nb = simple_notebook(images=code("fig", [error]))
    with pytest.raises(DeckError, match="raised ValueError"):
        read_notebook(write(tmp_path, nb))


def test_images_cell_with_no_output_is_refused(tmp_path):
    nb = simple_notebook(images=code("fig"))
    with pytest.raises(DeckError, match="produced no output"):
        read_notebook(write(tmp_path, nb))


def test_slide_needs_text_or_images(tmp_path):
    nb = notebook([nbformat.v4.new_markdown_cell("# Deck"),
                   nbformat.v4.new_markdown_cell("## Empty slide"),
                   nbformat.v4.new_markdown_cell("### notes\n\nonly notes")])
    with pytest.raises(DeckError, match="neither a '### text' nor an '### images'"):
        read_notebook(write(tmp_path, nb))


# --------------------------------------------------------------------------- markdown

def test_inline_runs_splits_emphasis_and_code():
    assert inline_runs("plain **bold** and *italic* and `code`") == [
        ("plain ", False, False), ("bold", True, False), (" and ", False, False),
        ("italic", False, True), (" and ", False, False), ("code", False, False)]


def test_blank_line_closes_a_bullet():
    """A lead-in after a list must not be merged into the last bullet."""
    blocks = [("markdown", "- first\n- second\n\n**Lead in** - and more\n\n- third", None)]
    items = parse_bullets(blocks)
    assert [("".join(r[0] for r in runs), is_bullet) for runs, _, is_bullet in items] == [
        ("first", True), ("second", True), ("Lead in - and more", False), ("third", True)]


def test_emphasis_wrapping_a_line_break_is_parsed():
    """Continuations are joined before inline markup is parsed."""
    blocks = [("markdown", "- the rate among pairs *close\n  enough to count*, not the rest", None)]
    (runs, _, _), = parse_bullets(blocks)
    assert ("close enough to count", False, True) in runs
    assert "*" not in "".join(run[0] for run in runs)


def test_indent_sets_level():
    blocks = [("markdown", "- top\n  - nested\n    - deeper", None)]
    assert [level for _, level, _ in parse_bullets(blocks)] == [0, 1, 2]


def test_estimate_height_grows_with_text():
    short = [([("a short bullet", False, False)], 0, True)]
    long = [([("a much longer bullet " * 8, False, False)], 0, True)]
    assert estimate_height(long, 4.0, 17) > estimate_height(short, 4.0, 17)


# --------------------------------------------------------------------------- building

def test_builds_a_deck_with_figure_table_and_notes(tmp_path):
    frame = pd.DataFrame({"count": ["1,110,640"]}, index=["step"])
    frame.index.name = "stage"
    nb = notebook([
        nbformat.v4.new_markdown_cell("# My deck"),
        nbformat.v4.new_markdown_cell("## Figure slide"),
        nbformat.v4.new_markdown_cell("### text\n\n- bullet one"),
        nbformat.v4.new_markdown_cell("### images"),
        code("fig", [display({"image/png": png_bytes()})]),
        nbformat.v4.new_markdown_cell("### notes\n\nsome notes"),
        nbformat.v4.new_markdown_cell("## Table slide"),
        nbformat.v4.new_markdown_cell("### text\n\n- bullet two"),
        nbformat.v4.new_markdown_cell("### images"),
        code("table", [display(frame_html(frame))]),
    ])
    out = build_deck(write(tmp_path, nb), out=tmp_path / "deck.pptx")

    from pptx import Presentation
    prs = Presentation(str(out))
    assert len(prs.slides) == 3                                   # section slide plus two
    assert prs.slides[1].shapes.title.text == "Figure slide"
    assert any(shape.shape_type == 13 for shape in prs.slides[1].shapes)       # a picture
    assert prs.slides[1].notes_slide.notes_text_frame.text == "some notes"
    tables = [shape for shape in prs.slides[2].shapes if shape.has_table]
    assert len(tables) == 1
    cells = [cell.text for cell in tables[0].table.rows[1].cells]
    assert "1,110,640" in cells          # the notebook's own formatting survives read_html


def test_multiindex_headers_are_flattened(tmp_path):
    frame = pd.DataFrame({"A": [1], "B": [2]}, index=["row"])
    frame.index.name = "stage"
    nb = notebook([
        nbformat.v4.new_markdown_cell("# Deck"),
        nbformat.v4.new_markdown_cell("## Slide"),
        nbformat.v4.new_markdown_cell("### text\n\n- bullet"),
        nbformat.v4.new_markdown_cell("### images"),
        code("table", [display(frame_html(frame))]),
    ])
    out = build_deck(write(tmp_path, nb), out=tmp_path / "deck.pptx")
    from pptx import Presentation
    table = next(shape for shape in Presentation(str(out)).slides[1].shapes
                 if shape.has_table).table
    headers = [cell.text for cell in table.rows[0].cells]
    assert headers == ["stage", "A", "B"]
    assert not any("Unnamed" in header for header in headers)


def test_section_slide_can_be_omitted(tmp_path):
    nb = simple_notebook(images=None)
    out = build_deck(write(tmp_path, nb), out=tmp_path / "deck.pptx", section_slide=False)
    from pptx import Presentation
    assert len(Presentation(str(out)).slides) == 1


def test_too_many_figures_is_refused(tmp_path):
    figures = [display({"image/png": png_bytes()}) for _ in range(3)]
    nb = simple_notebook(images=code("figs", figures))
    with pytest.raises(DeckError, match="3 figures on one slide"):
        build_deck(write(tmp_path, nb), out=tmp_path / "deck.pptx")


def test_table_and_figure_together_is_refused(tmp_path):
    frame = pd.DataFrame({"a": [1]})
    nb = simple_notebook(images=code("both", [display({"image/png": png_bytes()}),
                                              display(frame_html(frame))]))
    with pytest.raises(DeckError, match="a table and a figure on one slide"):
        build_deck(write(tmp_path, nb), out=tmp_path / "deck.pptx")


def test_overlong_table_is_refused(tmp_path):
    frame = pd.DataFrame({"value": range(40)})
    nb = simple_notebook(images=code("table", [display(frame_html(frame))]))
    with pytest.raises(DeckError, match="off the bottom of the slide"):
        build_deck(write(tmp_path, nb), out=tmp_path / "deck.pptx")


def test_overflowing_text_is_refused(tmp_path):
    bullets = "\n".join(f"- {'a long bullet that will not fit ' * 4} number {i}" for i in range(20))
    nb = simple_notebook(images=None)
    nb.cells[3].source = f"### text\n\n{bullets}"
    with pytest.raises(DeckError, match="overflows its placeholder"):
        build_deck(write(tmp_path, nb), out=tmp_path / "deck.pptx")


def test_missing_layout_is_refused(tmp_path):
    nb = simple_notebook(images=None)
    style = DeckStyle(content_layout="No Such Layout")
    with pytest.raises(DeckError, match="no layout named"):
        build_deck(write(tmp_path, nb), out=tmp_path / "deck.pptx", style=style)


# --------------------------------------------------------------------------- the shipped example

def test_shipped_example_builds(tmp_path):
    """`examples/minimal_deck.ipynb` must stay buildable, and stay executed in git.

    It is the documentation for the notebook format, so it rots the moment a rule changes without
    it. Building it here means a change that breaks it breaks the suite.
    """
    from pathlib import Path

    example = Path(__file__).resolve().parents[1] / "examples" / "minimal_deck.ipynb"
    if not example.exists():                       # a source checkout without the examples folder
        pytest.skip("examples/minimal_deck.ipynb is not present")

    out = build_deck(example, out=tmp_path / "example.pptx")
    from pptx import Presentation
    prs = Presentation(str(out))
    assert len(prs.slides) == 5                    # section slide plus the four it documents
    pictures = sum(1 for slide in prs.slides for shape in slide.shapes if shape.shape_type == 13)
    tables = sum(1 for slide in prs.slides for shape in slide.shapes if shape.has_table)
    assert (pictures, tables) == (3, 1)            # one figure, two stacked, one table
